# -*- coding: utf-8 -*-
"""按用户指示移除"用不了"的三类（②表征对齐、④元学习、⑥测试时适应），
连带：类别重编号为四大类、两表删行、TPT 移出复现清单、
参考文献删除 [19][20][21] 并将 [22]-[36] 顺延为 [19]-[33]、全文引用号同步重编。
"""
import copy
import re

import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

PATH = r"C:\Users\felix\Desktop\文献综述v8_修正版.docx"
d = docx.Document(PATH)
paras = list(d.paragraphs)
part = paras[0].part

# ---- 参考文献号 -> URL（按当前编号）----
REF_URL = {}
for p in paras:
    m = re.match(r"\[(\d+)\]\s", p.text)
    if m:
        hl = p._p.findall(qn("w:hyperlink"))
        if hl:
            rid = hl[0].get(qn("r:id"))
            if rid and rid in part.rels:
                REF_URL[int(m.group(1))] = part.rels[rid].target_ref
assert len(REF_URL) == 36, len(REF_URL)

DROP = {19, 20, 21}  # TENT / TPT / TT-SaD
renum = lambda n: n if n < 19 else n - 3
REF_URL2 = {renum(n): u for n, u in REF_URL.items() if n not in DROP}

LINK_RPR = ('<w:rPr><w:color w:val="0563C1"/><w:u w:val="single"/>'
            '<w:lang w:eastAsia="zh-CN"/></w:rPr>')
base_rPr = None
for r in paras[10]._p.findall(qn("w:r")):
    rpr = r.find(qn("w:rPr"))
    if rpr is not None:
        base_rPr = copy.deepcopy(rpr)
        break


def get_rid(url):
    for rid, rel in part.rels.items():
        if rel.is_external and rel.target_ref == url:
            return rid
    return part.relate_to(url, RT.HYPERLINK, is_external=True)


def esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def rebuild(p, text):
    pPr = p._p.find(qn("w:pPr"))
    for child in list(p._p):
        if child is not pPr:
            p._p.remove(child)
    for seg in re.split(r"(\[\d+\])", text):
        if not seg:
            continue
        m = re.fullmatch(r"\[(\d+)\]", seg)
        n = int(m.group(1)) if m else None
        if n in REF_URL2:
            rid = get_rid(REF_URL2[n])
            p._p.append(parse_xml(
                f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid}">'
                f'<w:r>{LINK_RPR}<w:t xml:space="preserve">{esc(seg)}</w:t></w:r>'
                f"</w:hyperlink>"))
        else:
            r = OxmlElement("w:r")
            if base_rPr is not None:
                r.append(copy.deepcopy(base_rPr))
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = seg
            r.append(t)
            p._p.append(r)


def find_para(prefix):
    for p in d.paragraphs:
        if p.text.startswith(prefix):
            return p
    raise KeyError(prefix)


# ============ 1. 删除三类段落 ============
for pref in ("（二）表征侧", "（四）学习策略侧", "（六）测试时侧"):
    p = find_para(pref)
    p._p.getparent().remove(p._p)

# ============ 2. 留下类别重编号 + 内部引用修正 ============
p = find_para("（三）优化侧")
rebuild(p, p.text.replace("（三）优化侧", "（二）优化侧"))
p = find_para("（五）先验侧")
rebuild(p, p.text.replace("（五）先验侧", "（三）先验侧"))
p = find_para("（七）条件组装侧")
rebuild(p, (p.text
            .replace("（七）条件组装侧", "（四）条件组装侧")
            .replace("⑦与⑤的组合", "④与③的组合")))

# ============ 3. 二节开篇：七大类 -> 四大类 ============
rebuild(find_para("Zhou 等[1]"), (
    "Zhou 等[1]、Wang 等[4]以及 Jahanifar 等[6]的综述确立了域泛化的分类学骨架。参照其框架，并结合各方法"
    "在生成管线中的干预位置与“可为本任务所用”的纳入标准，可将相关方法归为四大类，（一）至（四）依次对应"
    "数据、优化、先验与条件组装四个侧面。下文每类均按统一框架展开：先界定该类在“三源域→一目标域”场景中"
    "的干预位置与所回应的矛盾，再逐篇给出库内代表文献的核心思想，最后归纳文献口径的优缺点与本课题的取舍。"))

# ============ 4. 检索范围段：写入纳入标准，36 -> 33 ============
rebuild(find_para("在文献检索上"), (
    "在文献检索上，本综述的范围与流程明确如下。范围以“虚拟染色（virtual staining / virtual histology "
    "staining）× 域泛化（domain generalization）”为主轴；就检索范围所及，未检索到以二者交叉为主题的专门"
    "综述，交叉证据散见于虚拟染色综述对跨站点鲁棒性的讨论[15][16]与计算病理域泛化综述[6]。鉴于直接交叉的"
    "文献体量有限，按“问题本质相同即可借鉴”的原则放宽一圈纳入：计算病理与医学影像域泛化（其 train–test "
    "差异体现为染色、扫描仪与中心偏移[5][7]），以及通用域泛化中思想可迁移的方法——RGB→红外、MR→CT 等"
    "跨模态翻译与跨器官染色共享“训练与测试分布存在差异”的问题本质，仅差异的表现轴不同。纳入标准为“思想"
    "可为本任务所用”：测试时适应在推理期利用目标域输入做适应，与训练期零接触的域泛化是两个问题设定；源域"
    "特征对齐与元学习经系统性评测判定收益不稳定[7][8]——以上三类均不纳入本综述。流程上，以三篇域泛化综述"
    "[1][4][6]与病理域泛化基准[7]的参考文献列表为起点滚雪球扩展，检索工具仅用于生成候选，最终收录的 33 条"
    "文献均经逐篇核对原文后进入本综述；另以每日文献雷达持续巡查新文献并做撞车检测，候选文献经核验后滚动"
    "补入对应类别。"))

# ============ 5. （八）注 + 统计段 ============
rebuild(find_para("注：此处统计"), (
    "注：此处统计为各类别核心代表作篇数，骨干模型与综述类文献未重复计入；③类仅统计“基础模型用于域泛化”"
    "的文献，基础模型本身（UNI/CONCH/TITAN）另计为工具背景，不入方法类。"))
rebuild(find_para("为确定研究路线"), (
    "为确定研究路线的落点，统计本综述收录文献（33 篇）在各类别的分布与时间走向：数据侧历史存量最大（增广"
    "是所有配方的默认基础），并在 2024 年后出现“扩散合成域（源域多样性不够，就用扩散模型凭空生成更多未知"
    "域的训练数据）”的新支线[11]；优化侧保持稳定的小体量（SWAD[18]/ERM++[9]一脉）；先验侧在 2024–2026 年"
    "集中爆发——FLEX[25]、VFM-DG[26]、UNIStainNet[29]密集出现，是当前最活跃的类别，同期基础模型本身"
    "（UNI[22]、CONCH[23]、TITAN[24]）也密集发布，巨量多中心无标注数据预训练的特征自带跨域鲁棒性，构成"
    "该路线的燃料；检索条件化在 2017 年奠基[2][3]后沉寂，正随检索增强思想的整体回潮而复兴。近三年文献明确"
    "集中于“①扩散合成域”与“③基础模型先验”。"))

# ============ 6. 表0 / 表1：删行 + 重编号 ============
def fix_table(tb, del_keys, rename):
    for row in list(tb.rows):
        key = row.cells[0].text.strip()
        if key in del_keys:
            tb._tbl.remove(row._tr)
        elif key in rename:
            row.cells[0].text = rename[key]

fix_table(d.tables[0],
          {"② 表征侧", "④ 元学习", "⑥ 测试时"},
          {"③ 优化侧": "② 优化侧",
           "⑤ 先验侧（基础模型的DG利用）": "③ 先验侧（基础模型的DG利用）",
           "⑦ 检索条件化": "④ 检索条件化"})
fix_table(d.tables[1],
          {"② 表征侧：对齐/解耦", "④ 元学习", "⑥ 测试时适应"},
          {"③ 优化侧：平坦最小值": "② 优化侧：平坦最小值",
           "⑤ 先验侧：基础模型注入": "③ 先验侧：基础模型注入",
           "⑦ 检索条件化": "④ 检索条件化"})

# ============ 7. （九）引句与排除链 ============
rebuild(find_para("分类与趋势回答了"), (
    "分类与趋势回答了“领域在做什么”，选题还须回答“哪条路线适合本任务”。以下以四条判定轴检验四大类方法："
    "A 轴——是否需要目标域任何数据（训练期目标器官的 H&E 与无标注输入均不可见）；B 轴——生成保真风险"
    "（细胞级结构不可损伤）；C 轴——计算与工程成本；D 轴——病理语义风险（合成过程不得扭曲诊断语义）。"
    "判定结果汇总如下："))
rebuild(find_para("上述判定构成"), (
    "上述判定构成一条“排除—选择”链：数据增广无法解决形态泛化，只能作为基础保留；优化侧以稳定优化提供零"
    "语义风险的增益，作基础叠加；两个常见的候选方案被 A 轴否决——以目标域 H&E 统计量作染色归一化锚点、"
    "源域掺入目标域无标注输入的混合训练，均直接违反 A 轴约束。真正同时通过四轴检验的主路线只有 ③×④ 组合："
    "基础模型注入回应“认得出是什么组织”（形态学泛化），源域检索条件化回应“该染成什么样”（逐图风格锚定），"
    "恰好分别覆盖引言所述核心矛盾的两侧。"))

# ============ 8. 四节：复现清单去掉 TPT ============
rebuild(find_para("（二）各大类代表"), (
    "（二）各大类代表（3 项）：数据侧复现 RandStainNA[17]——病理特异的染色增广，重点对照维度：叠加于骨干"
    "后对跨器官指标的边际增益，直接检验“增广能否覆盖形态差异”的命题；优化侧复现 SWAD——平坦最小值代表"
    "（已有既定复现计划），重点对照维度：零语义风险的增益在生成任务上的可叠加性；先验侧复现 UNIStainNet"
    "[29]——本路线最近邻，重点对照维度：SPADE 与 Zero-Conv 两种注入接口之差、以及是否利用检索条件。"))
p = find_para("选择理由归纳为三条")
rebuild(p, p.text.replace("上述十项", "上述九项"))

# ============ 9. 五节与结语：清除对已删类的提及 ============
rebuild(find_para("综上，面向"), (
    "综上，面向“训练期目标器官完全不可见的跨器官 PI→H&E 虚拟染色”，现有方法存在三个具体缺口：其一，形态"
    "学泛化无解——风格层面的增广不覆盖跨器官的组织形态差异（第二节（九）判定）；其二，先验用法单一——基础"
    "模型至多以单路径条件进入生成通路，无“结构 + 风格”双先验的先例（第三节（二））；其三，验证协议宽松——"
    "严格留出器官零样本评测缺失，合规的风格条件化变体（源域库自查询）亦未被探索（第三节（四））。"))
rebuild(find_para("这条路线的必要性"), (
    "这条路线的必要性由排除法支撑：数据增广只能覆盖风格维度，作基础而非主路线；在其余合规选项中，唯有"
    "“基础模型双先验注入 × 检索条件化”同时回应形态学泛化与风格条件化两个侧面，且近邻六篇核查确认该组合"
    "在虚拟染色任务上无先例。"))
rebuild(find_para("本综述以"), (
    "本综述以“训练期目标器官完全不可见”的跨器官 PI→H&E 虚拟染色为对象，完成了从分类图景、趋势统计、适用"
    "性判定到重点方向微观差异化的完整论证：数据侧与优化侧只消费源域数据，宜作配方基础；基础模型冻结注入与"
    "源域检索条件化的组合，在合规性之上同时覆盖形态学泛化与风格条件化，且经近邻核查确认无先例。据此，“冻结"
    "病理基础模型提供结构与风格先验、桥式扩散骨干负责生成、源域检索库实现逐图条件化”构成本课题的落点。文献"
    "综述并非一次性工作，后续将随实验进展持续补充近邻文献并更新本综述。"))

# ============ 10. 全文引用号重编（正文段落，[22]-[36] -> [19]-[33]）============
# 安全检查：[19][20][21]（TENT/TPT/TT-SaD）应已无正文引用
_leftover = [(i, p.text[:40]) for i, p in enumerate(d.paragraphs)
             if re.search(r"\[(19|20|21)\]", p.text) and not re.match(r"\[\d+\]\s", p.text)]
assert not _leftover, _leftover

ref_heading_seen = False
for p in d.paragraphs:
    if p.style.name.startswith("Heading") and p.text.startswith("参考文献"):
        ref_heading_seen = True
        continue
    if ref_heading_seen:
        continue  # 文献条目单独处理
    toks = [int(n) for n in re.findall(r"\[(\d+)\]", p.text)]
    if any(n >= 19 for n in toks):
        new_text = re.sub(r"\[(\d+)\]",
                          lambda m: f"[{renum(int(m.group(1)))}]"
                          if int(m.group(1)) >= 19 else m.group(0), p.text)
        rebuild(p, new_text)

# 表格单元格（无超链接，run 级替换）
for tb in d.tables:
    for row in tb.rows:
        for cell in row.cells:
            if re.search(r"\[(1[9]|2\d|3\d)\]", cell.text):
                for t in cell._tc.iter(qn("w:t")):
                    if t.text and re.search(r"\[\d+\]", t.text):
                        t.text = re.sub(
                            r"\[(\d+)\]",
                            lambda m: f"[{renum(int(m.group(1)))}]"
                            if int(m.group(1)) >= 19 else m.group(0), t.text)

# ============ 11. 参考文献条目：删 [19][20][21]，顺延重编号 ============
for p in list(d.paragraphs):
    m = re.match(r"\[(\d+)\]\s", p.text)
    if not m:
        continue
    old = int(m.group(1))
    if old in DROP:
        p._p.getparent().remove(p._p)
    elif old > 21:
        for t in p._p.iter(qn("w:t")):
            if t.text and t.text.startswith(f"[{old}]"):
                t.text = f"[{old-3}]" + t.text[len(f"[{old}]"):]
                break

d.save(PATH)
print("saved:", PATH)
