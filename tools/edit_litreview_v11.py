# -*- coding: utf-8 -*-
"""（四）类段落改写为可懂表述，类名 条件组装侧 -> 条件侧；同步二节开篇类别列表。"""
import copy
import re

import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

PATH = r"C:\Users\felix\Desktop\文献综述v8_修正版.docx"
d = docx.Document(PATH)
part = d.paragraphs[0].part

REF_URL = {}
for p in d.paragraphs:
    m = re.match(r"\[(\d+)\]\s", p.text)
    if m:
        hl = p._p.findall(qn("w:hyperlink"))
        if hl:
            rid = hl[0].get(qn("r:id"))
            if rid and rid in part.rels:
                REF_URL[int(m.group(1))] = part.rels[rid].target_ref

LINK_RPR = ('<w:rPr><w:color w:val="0563C1"/><w:u w:val="single"/>'
            '<w:lang w:eastAsia="zh-CN"/></w:rPr>')
base_rPr = None
for p in d.paragraphs:
    if p.text.startswith("（一）数据侧"):
        for r in p._p.findall(qn("w:r")):
            rpr = r.find(qn("w:rPr"))
            if rpr is not None:
                base_rPr = copy.deepcopy(rpr)
                break
        break


def get_rid(url):
    for rid, rel in part.rels.items():
        if rel.is_external and rel.target_ref == url:
            return rid
    return part.relate_to(url, RT.HYPERLINK, is_external=True)


def rebuild(p, text):
    pPr = p._p.find(qn("w:pPr"))
    for child in list(p._p):
        if child is not pPr:
            p._p.remove(child)
    for seg in re.split(r"(\[\d+\])", text):
        if not seg:
            continue
        m = re.fullmatch(r"\[(\d+)\]", seg)
        if m and int(m.group(1)) in REF_URL:
            rid = get_rid(REF_URL[int(m.group(1))])
            p._p.append(parse_xml(
                f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid}">'
                f'<w:r>{LINK_RPR}<w:t xml:space="preserve">{seg}</w:t></w:r>'
                f"</w:hyperlink>"))
        else:
            r = OxmlElement("w:r")
            if base_rPr is not None:
                r.append(copy.deepcopy(base_rPr))
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = seg
            r.append(t)
            p._p.append(r)


def find_para(prefix):
    for p in d.paragraphs:
        if p.text.startswith(prefix):
            return p
    raise KeyError(prefix)


rebuild(find_para("（四）条件组装侧"), (
    "（四）条件侧：检索与范例条件化。生成模型输出一张染色图像需要两类信息：画什么结构（来自输入图像本身），"
    "以及照什么风格去染（生成条件）。前三类方法分别改数据、改优化、改注入的先验，这一类改的是“风格条件从"
    "哪里来”：不为网络学一个全局统一的风格映射，而是在生成每一张图时临时找参考图，告诉模型“照这个样子染”，"
    "逐图锚定输出风格——回应的正是引言所述“风格条件不能来自对目标域的猜测”这一矛盾。库内的范例式文献确立了"
    "“以参考图定输出风格”的范式：范例着色[2]以用户交互的参考图引导黑白图像上色；深度类比[3]以 a 图的结构×"
    "b 图的风格合成新图；病理中最常用的染色归一化（Beer-Lambert 一脉）同样以一张标准 H&E 参考图为颜色之锚。"
    "但范例式的共同前提是“手头有目标域的参考图”——归一化的锚图取自目标域，范例着色的参考由用户提供——在"
    "“目标域 H&E 训练期不可见”的约束下直接违规，且人工交互难以批量量化。可用的合规变体是自查询式：参考图不"
    "从目标域拿，而以测试输入自身为查询，在仅由源域构建的检索库中取最近邻——它锚定的不是“目标域该长什么样”，"
    "而是“源域中最像它的样本被染成了什么样”，既不接触目标域，又逐图个性化、不猜平均风格（详见第三节（三））。"
    "对本课题的启发：范例式提供思想（以参考图定风格），自查询式提供合规入口（源域库 + 测试输入查询），④与③"
    "的组合构成本课题主路线的风格侧。"))

p = find_para("Zhou 等[1]")
rebuild(p, p.text.replace("数据、优化、先验与条件组装四个侧面", "数据、优化、先验与条件四个侧面"))

d.save(PATH)
print("saved")
